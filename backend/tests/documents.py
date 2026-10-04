"""Build small but real PDF and Word files for upload tests."""

import io
import zipfile

import docx
from pypdf import PdfWriter


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(lines: list[str], pages: int = 1) -> bytes:
    """A valid PDF with the given lines of Helvetica text on every page."""
    shown = " ".join(f"({_escape(line)}) Tj T*" for line in lines)
    content = f"BT /F1 12 Tf 14 TL 72 720 Td {shown} ET"
    page_ids = [4 + 2 * i for i in range(pages)]
    objects: dict[int, str] = {
        1: "<< /Type /Catalog /Pages 2 0 R >>",
        2: f"<< /Type /Pages /Kids [{' '.join(f'{i} 0 R' for i in page_ids)}] /Count {pages} >>",
        3: "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for page_id in page_ids:
        objects[page_id] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_id + 1} 0 R >>"
        )
        objects[page_id + 1] = f"<< /Length {len(content)} >>\nstream\n{content}\nendstream"

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n{objects[number]}\nendobj\n".encode("latin-1")
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for number in sorted(objects):
        out += f"{offsets[number]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode()
    out += f"startxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(out)


def encrypt_pdf(pdf: bytes, password: str = "secret") -> bytes:
    writer = PdfWriter(clone_from=io.BytesIO(pdf))
    writer.encrypt(password)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_zip_bomb(unzipped_bytes: int) -> bytes:
    """A tiny .docx-shaped zip that would inflate to ``unzipped_bytes``."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * unzipped_bytes)
    return out.getvalue()
